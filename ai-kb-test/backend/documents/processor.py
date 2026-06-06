import re
from pathlib import Path
from typing import List, Tuple
from config import CHUNK_SIZE, CHUNK_OVERLAP
from database.sqlite_db import add_chunk

def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    mapping = {'.pdf': 'pdf', '.docx': 'word', '.doc': 'word', '.xlsx': 'excel', '.xls': 'excel', '.pptx': 'ppt', '.ppt': 'ppt', '.txt': 'text', '.md': 'text'}
    return mapping.get(ext, 'unknown')

def extract_text(file_path: str, doc_type: str) -> str:
    if doc_type == 'text':
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except:
            with open(file_path, 'r', encoding='gbk') as f:
                return f.read()
    elif doc_type == 'word':
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            return f"[Word解析失败: {e}]"
    elif doc_type == 'pdf':
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            return text
        except Exception as e:
            return f"[PDF解析失败: {e}]"
    elif doc_type == 'excel':
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            parts = []
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(v) if v is not None else "" for v in row)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    parts.append(f"--- Sheet: {sheet.title} ---\n" + "\n".join(rows))
            return "\n\n".join(parts)
        except Exception as e:
            return f"[Excel解析失败: {e}]"
    else:
        return f"[不支持的文档类型: {doc_type}]"

def semantic_chunk(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[Tuple[int, str]]:
    paragraphs = re.split(r'\n\s*\n', text.strip())
    chunks = []
    current_chunk = []
    current_len = 0
    chunk_idx = 0
    for para in paragraphs:
        para = para.strip()
        if not para: continue
        para_len = len(para)
        if para_len > chunk_size:
            sentences = re.split(r'(?<=[。！？.!?])\s+', para)
            for sent in sentences:
                if current_len + len(sent) > chunk_size and current_chunk:
                    chunks.append((chunk_idx, "\n".join(current_chunk)))
                    chunk_idx += 1
                    overlap_text = "\n".join(current_chunk)
                    overlap_chars = overlap_text[-overlap:] if len(overlap_text) > overlap else overlap_text
                    current_chunk = [overlap_chars, sent]
                    current_len = len(overlap_chars) + len(sent)
                else:
                    current_chunk.append(sent)
                    current_len += len(sent)
        else:
            if current_len + para_len > chunk_size and current_chunk:
                chunks.append((chunk_idx, "\n".join(current_chunk)))
                chunk_idx += 1
                overlap_text = "\n".join(current_chunk)
                overlap_chars = overlap_text[-overlap:] if len(overlap_text) > overlap else overlap_text
                current_chunk = [overlap_chars, para]
                current_len = len(overlap_chars) + para_len
            else:
                current_chunk.append(para)
                current_len += para_len
    if current_chunk:
        chunks.append((chunk_idx, "\n".join(current_chunk)))
    return chunks

def process_document(file_path: str, doc_id: int, filename: str) -> int:
    doc_type = detect_file_type(filename)
    text = extract_text(file_path, doc_type)
    if not text.strip():
        return 0
    chunks = semantic_chunk(text)
    for idx, chunk_text in chunks:
        add_chunk(doc_id, idx, chunk_text)
    return len(chunks)
